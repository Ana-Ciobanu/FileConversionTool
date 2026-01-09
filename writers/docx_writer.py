from docx import Document
import logging
from writers.base_writer import FileWriter


class DOCXWriter(FileWriter):
    """
    Writes text or table data to a .docx file.
    Supports data from TXT, DOCX, PDF (as text), and CSV (as table).
    """

    supported_input_types = {"text", "table"}

    def write(self, data: dict, output_path: str):
        try:
            doc = Document()
            if data.get("type") == "text":
                content = data.get("content", "")
                for line in content.split("\n"):
                    doc.add_paragraph(line)
            elif data.get("type") == "table":
                rows = data.get("rows", [])
                if not rows:
                    raise ValueError("No rows provided for table data")
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        table.cell(i, j).text = str(cell)
            else:
                raise ValueError("DOCXWriter supports only text or table data")
            doc.save(output_path)
        except Exception as e:
            logging.error(f"Failed to write DOCX file {output_path}: {e}")
            raise
